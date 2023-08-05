# pdfminer or pdfminer.six
import pdfminer
#!c:\bin\python27\python.exe
import sys
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.pdfdevice import PDFDevice, TagExtractor
from pdfminer.pdfpage import PDFPage
from pdfminer.converter import XMLConverter, HTMLConverter, TextConverter, PDFLayoutAnalyzer
from pdfminer.cmapdb import CMapDB
from pdfminer.layout import LAParams,LTContainer,LTText,LTTextBox,LTFigure,LTImage,LTTextLineHorizontal
import re
import collections
import bisect
import functools
import six

MARGIN = 1
MIN_INTERSECT = 4
PAD = ' '
CHAR_SIZE = 4.7
RE_LONG_SPACES = re.compile('( {2,})')

@functools.total_ordering
class Interval(object):
    __slots__ = ['a', 'b', 'margin']

    def __init__(self,a,b,margin=MARGIN):
        self.a = a
        self.b = b
        self.margin = margin

    def __or__(self,o):
        if not (o & self)  : return None
        return Interval(min(self.a,o.a),max(self.b,o.b))

    def __and__(self,o):
        a1 = max(self.a,o.a)
        b1 = min(self.b,o.b)
        if b1>a1:
            return Interval(a1,b1)
        else:
            return None

    def hull(self,o):
        return Interval(min(self.a,o.a),max(self.b,o.b))

    def dist(self,o):
        x,y = sorted([self,o])
        return max(0,y.a-x.b)

    @property
    def size(self):
        return self.b - self.a

    def __contains__(self,o):
        return self.b + self.margin >= o.b >= o.a >= self.a - self.margin

    def __gt__(self,o):
        return self.a>o.b

    def __eq__(self,o):
        return (self & o) is not None

    def __str__(self):
        return "[%s - %s]"%(self.a,self.b)

    def __repr__(self):
        return "<Interval(%s,%s)>"%(self.a,self.b)

class Position:
    def __init__(self,x0,y0,x1,y1):
        self.x0,self.y0,self.x1,self.y1 = x0,y0,x1,y1
        self.h = self.y1 - self.y0
        self.w = self.x1 - self.x0
        self.y_int = Interval(y0,y1,MARGIN*self.h)
        self.x_int = Interval(x0,x1,MARGIN*self.w)

    def __str__(self):
        return " x,y = ( %.0f, %.0f ) w,h = ( %.0f, %.0f )"%(self.x0,self.y0,self.w,self.h)

    @property
    def rightx(self):
        return self.x1

    def label(self):
        return "(%.0f, %.0f)"%(self.x0,self.y0)
    
    def center(self):
        return self.x0+self.w/2, self.y0 + self.h/2

    def same_y(self,o):
        return o.y_int in self.y_int or self.y_int in o.y_int

    

@functools.total_ordering
class TextFrame(object):
    char_size = CHAR_SIZE

    def __init__(self,position,text):
        self.position = position
        self.text = text

    def split_vertical(self):
        lines = self.text.split('\n')
        line_height = self.position.h/len(lines)
        for i,t in enumerate(lines):
            tpos = Position(self.position.x0,self.position.y1-(i+1)*line_height,self.position.x1,self.position.y1-i*line_height)
            pad = int(self.position.w/self.char_size)-len(t)
            if pad <0: pad =0
            yield TextFrame(tpos,t+PAD*pad)

    def split_horizontal(self):
        if not self.text:
            yield self.text
            return
        char_size = self.position.w/len(self.text)
        x0 = self.position.x0
        sub_text = iter(RE_LONG_SPACES.split(self.text))
        for t in sub_text:
            x1 = x0 + char_size * len(t)
            tpos = Position(x0,self.position.y0,self.position.x1,self.position.y1)
            yield TextFrame(tpos,t)
            x0 = x1 + char_size * len(next(sub_text)) # for spaces

    def __eq__(self,o):
        return self.position.x0 == o.position.x0

    def __gt__(self,o):
        return self.position.x0 > o.position.x0

    def __repr__(self):
        return "<TF %s: %s>"%(self.position,repr(self.text))

class AlignedTextFrame(object):
    def __init__(self,y_int,frames=[]):
        self.interval = y_int
        self.frames = []
        self.frames.extend(frames)

    def __repr__(self):
        return repr(u"%s: %s"%(self.interval,",".join(i.text for i in self.frames)))

    def is_aligned(self,text_frame):
        intersect = text_frame.position.y_int & self.interval
        if intersect is None: return 0
        return intersect.size
        #return intersect is not None and len(intersect) > MIN_INTERSECT

    def add_frame(self,text_frame):
        #self.interval = text_frame.position.y_int | self.interval
        self.frames.extend(text_frame.split_horizontal())

    def merge(self,aligned):
        new_interval = self.interval | aligned.interval
        if new_interval:
            self.interval = new_interval
            self.frames.extend(aligned.frames)

#    def __cmp__(self,o):
#        if isinstance(o,AlignedTextFrame):
#            return cmp(self.interval,o.interval)
#        if isinstance(o,TextFrame):
#            return cmp(self.interval,o.position.y_int)
#        return cmp(self.interval,o)

    def get_padded_texts(self):
        prev_pos_x = 0
        result = []
        for tf in sorted(self.frames):
            result.append(PAD*int((tf.position.x0 - prev_pos_x)/tf.char_size))
            prev_pos_x = tf.position.x1
            result.append(tf.text)
        return  ''.join(result)


class Page:
    def __init__(self):
        self.aligned_frames =[]

    def add_frame(self,frame):
        if not frame.text.strip():
            return
        candidates = []
        insert = False
        if not self.aligned_frames:
            self.aligned_frames.append(AlignedTextFrame(frame.position.y_int,[frame]))
            return
        for i in range(len(self.aligned_frames)):
            line = self.aligned_frames[i]
            alignment =  line.is_aligned(frame)
            if alignment:
                candidates.append((alignment,line))
            if frame.position.y_int>line.interval:
                insert = True
                break
        if not candidates:
            if insert:
                self.aligned_frames.insert(i,AlignedTextFrame(frame.position.y_int,[frame]))
            else:
                self.aligned_frames.append(AlignedTextFrame(frame.position.y_int,[frame]))    
        else:
            _,line = sorted(candidates)[-1]
            line.add_frame(frame)

    def add_text(self,text_frame):
        for tf in text_frame.split_vertical():
            self.add_frame(tf)

    def setMediaBox(self,mediabox):
        self.x,self.y,self.w,self.h = mediabox


class PdfTable(object):
    def __init__(self,rows):
        self.rows = rows
        self._get_columns(self.rows) #list of positions

    def _get_columns(self,aligned_frames):
        self.columns = sorted(i.position.x_int for i in aligned_frames[0].frames)
        for row in aligned_frames[1:]:
            self._add_row(row.frames)
        
    def _add_row(self,row):
        for cell in row:
            ids = [i for i,col in enumerate(self.columns) if col & cell.position.x_int]
            if len(ids)==0: # new columnn
                bisect.insort(self.columns,cell.position.x_int)
            else:
                for i in ids:
                    self.columns[i] &= cell.position.x_int
        return len([1 for cell1,cell2 in zip(columns,columns[1:]) if cell1&cell2])==0

    def get_table(self):
        table = []
        for alignedframes in self.rows:
            row = ['']*len(self.columns)
            for frame in alignedframes.frames:
                for i,col in enumerate(self.columns):
                    if col & frame.position.x_int:
                        assert not row[i]
                        row[i]=frame.text
                        break
            table.append(row)
        return table

    def merge_margin(self,small_col=15,margin=5):
        result = []
        for col in self.columns:
            if result:
                pre = result[-1]
                if pre.dist(col)<margin:
                    if min(col.size,pre.size) < small_col:
                        result.pop(-1)
                        col = pre.hull(col)
            result.append(col)
        self.columns = result
        return self     

class TextAnalyzer(TextConverter):
    def __init__(self,*args,**kwargs):
        TextConverter.__init__(self,*args,**kwargs)
        self.pages = {}

    def render(self,item):
        if isinstance(item,LTFigure):
             return
        if isinstance(item, LTContainer):
            for child in item:
                self.render(child)
        if isinstance(item,LTTextLineHorizontal):
            self.pages.setdefault(self.pageno,Page()).add_text(TextFrame(Position( item.x0,item.y0,item.x1,item.y1), item.get_text()))

    def begin_page(self,page,ctm):
        result= TextConverter.begin_page(self,page,ctm)
        self.pages.setdefault(self.pageno+1,Page()).setMediaBox(page.mediabox)
        return result
    
    def end_page(self,page):
        #assert self.text=='',Exception(self.text)
        return TextConverter.end_page(self,page)

    def receive_layout(self, ltpage):
        if self.showpageno:
            self.write_text('Page %s\n' % ltpage.pageid)
        self.render(ltpage)
        #self.write_text('\f')
        return

    def all_text(self):
        for pdfpage,page in sorted(six.iteritems(self.pages)):
            yield [aligned_frame.get_padded_texts() for aligned_frame in page.aligned_frames]

    def get_result(self,page):
        return {pageno:PdfTable(page.aligned_frames).merge_margin() for pageno,page in six.iteritems(self.pages)}

def extract_text(f,*page_numbers):
    rsrcmgr = PDFResourceManager(caching=True)
    codec = 'utf-8'
    laparams = LAParams()
    laparams.all_texts = False
    device = TextAnalyzer(rsrcmgr, sys.stdout, codec=codec,laparams = laparams)
    read(f,device,page_numbers)
    device.close()
    return device.get_result()

if __name__ == '__main__':
    import sys
    import openpyxl
    filename = sys.argv[1]
    with file(filename,'rb') as f:
        pages= extract_text(f)
        for i in range(min(pages),max(pages)+1):            
            print (''.join(i).encode('utf-8'))
